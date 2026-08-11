#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Minimal Markdown -> .docx converter for the system-design doc."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "D:/MY-WEB/系统设计.md"
DST = "D:/MY-WEB/系统设计.docx"

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(8)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'F2F2F2')
    pPr.append(shd)
    for ln in text.split('\n'):
        r = p.add_run(ln)
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        # ensure east-asian font too
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        p.add_run('').add_break() if False else None
        # line break
        if ln != text.split('\n')[-1]:
            r.add_break()
    return p

def add_inline(p, text):
    # handle **bold**
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
        elif part:
            p.add_run(part)

def is_table_row(line):
    return line.strip().startswith('|') and line.strip().endswith('|')

def is_sep(line):
    s = line.strip().strip('|').replace(' ', '')
    return bool(re.fullmatch(r'-{1,}:?(-{1,}:?)*', s)) and '|' not in s

def parse_table(lines, i):
    # header at i, sep at i+1, body until non-table row
    header = [c.strip() for c in lines[i].strip().strip('|').split('|')]
    j = i + 2
    body = []
    while j < len(lines) and is_table_row(lines[j]) and lines[j].strip() != '':
        body.append([c.strip() for c in lines[j].strip().strip('|').split('|')])
        j += 1
    return header, body, j

def main():
    with open(SRC, encoding='utf-8') as f:
        lines = f.read().split('\n')
    doc = Document()
    # base font
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        if line.strip() == '':
            i += 1; continue
        if line.lstrip().startswith('```'):
            # code block
            lang = line.strip().strip('`').strip()
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1  # skip closing ```
            add_code_block(doc, '\n'.join(buf))
            continue
        if is_table_row(line) and i + 1 < n and is_sep(lines[i+1]):
            header, body, j = parse_table(lines, i)
            t = doc.add_table(rows=1, cols=len(header))
            t.style = 'Table Grid'
            hdr = t.rows[0].cells
            for k, h in enumerate(header):
                hdr[k].text = ''
                r = hdr[k].paragraphs[0].add_run(h); r.bold = True
                set_cell_bg(hdr[k], 'DCE6F1')
            for row in body:
                cells = t.add_row().cells
                for k, val in enumerate(row):
                    cells[k].text = ''
                    add_inline(cells[k].paragraphs[0], val)
            doc.add_paragraph()
            i = j; continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            p = doc.add_heading(level=level)
            add_inline(p, m.group(2))
            i += 1; continue
        if re.match(r'^-{2,}\s*$', line.strip()):
            i += 1; continue
        if re.match(r'^(\s*-\s+)', line):
            # bullet list
            items = []
            while i < n and re.match(r'^(\s*-\s+)', lines[i]):
                items.append(re.sub(r'^\s*-\s+', '', lines[i])); i += 1
            for it in items:
                p = doc.add_paragraph(style='List Bullet')
                add_inline(p, it)
            continue
        if re.match(r'^(\s*\d+\.\s+)', line):
            items = []
            while i < n and re.match(r'^(\s*\d+\.\s+)', lines[i]):
                items.append(re.sub(r'^\s*\d+\.\s+', '', lines[i])); i += 1
            for it in items:
                p = doc.add_paragraph(style='List Number')
                add_inline(p, it)
            continue
        # normal paragraph
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.save(DST)
    print("saved", DST)

if __name__ == '__main__':
    main()
